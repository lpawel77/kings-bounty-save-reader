"""Okienkowa przeglądarka zapisu King's Bounty: Armored Princess.

Pokazuje, co gra wylosowała w każdym budynku, skrzyni i armii na wszystkich
mapach (także nieodwiedzonych), oraz czy dany przedmiot w ogóle pojawi się w grze.

    python kb_app.py [plik.sav]
"""

import contextlib
import io
import json
import os
import sys
from collections import defaultdict
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from kb_gamedata import (GRUPY_PREMII, RASY, SLOTY, SZKOLY, GameData, bonus_good, bonus_text,
                         find_game_dirs, is_game_dir)

try:
    from PIL import Image, ImageTk
except ImportError:  # bez Pillow aplikacja działa, tylko bez obrazków
    Image = ImageTk = None
from kb_mapview import MapView
from kb_save import load_save, print_report
from kb_world import ARMY, ITEM, RES, ROLA, SPELL, UNIT, World, pretty_atom
from kb_world import _miejsc as miejsc

APP_TITLE = "King's Bounty – czytnik zapisu"
CONFIG = Path(os.environ.get("APPDATA") or Path.home()) / "kb_save_reader.json"

KINDS = [(ITEM, "Przedmioty"), (SPELL, "Zwoje i czary"), (UNIT, "Jednostki do werbunku"),
         (ARMY, "Armie"), (RES, "Zasoby (złoto, kryształy...)")]
TYPE_ORDER = ["Zamek", "Budynek / zamek", "Budynek / sklep", "Postać", "Skrzynia",
              "Ołtarz / skrytka", "Ołtarz", "Znalezisko", "Armia"]
STATUS_COLORS = {"masz": "#1a7f37", "jest": "", "ulepszenie": "#0b5cad", "byl": "#b35900",
                 "skrypt": "#7a4fb3", "brak": "#8a8a8a"}
STATUS_LABELS = [("", "Wszystkie"), ("masz", "Masz"), ("jest", "Do zdobycia"),
                 ("ulepszenie", "Tylko przez ulepszenie"), ("byl", "Był, ale już go nie ma"),
                 ("skrypt", "Tylko z zadań / wydarzeń"), ("brak", "Nie pojawi się")]
HIDDEN_SLOTS = ("medal", "hidden", "setbonus")
BUILDING_TYPES = ("Zamek", "Budynek / zamek", "Budynek / sklep", "Postać")


def save_dirs():
    docs = [Path.home() / "Documents"]
    if os.environ.get("OneDrive"):
        docs.append(Path(os.environ["OneDrive"]) / "Documents")
    for d in docs:
        for name in ("Kings Bounty Princess", "King's Bounty Armored Princess", "Kings Bounty Crossworlds"):
            p = d / "My Games" / name / "$save"
            if p.is_dir():
                yield p


def sort_key(value):
    v = str(value).replace(" ", "").replace(" ", "")
    try:
        return (0, float(v))
    except ValueError:
        return (1, str(value).casefold())


def make_sortable(tree, columns):
    def sort(col, desc):
        rows = [(tree.set(k, col), k) for k in tree.get_children("")]
        rows.sort(key=lambda r: sort_key(r[0]), reverse=desc)
        for i, (_, k) in enumerate(rows):
            tree.move(k, "", i)
        tree.heading(col, command=lambda: sort(col, not desc))
    for c in columns:
        tree.heading(c, command=lambda c=c: sort(c, False))


def scrolled(parent, widget_cls, **kw):
    frame = ttk.Frame(parent)
    w = widget_cls(frame, **kw)
    ys = ttk.Scrollbar(frame, orient="vertical", command=w.yview)
    w.configure(yscrollcommand=ys.set)
    w.grid(row=0, column=0, sticky="nsew")
    ys.grid(row=0, column=1, sticky="ns")
    if widget_cls is ttk.Treeview:
        xs = ttk.Scrollbar(frame, orient="horizontal", command=w.xview)
        w.configure(xscrollcommand=xs.set)
        xs.grid(row=1, column=0, sticky="ew")
    frame.rowconfigure(0, weight=1)
    frame.columnconfigure(0, weight=1)
    return frame, w


def table(parent, columns, height=10, icon=False):
    """Treeview z kolumnami [(id, nagłówek, szerokość, wyrównanie)] i sortowaniem.

    icon=True dodaje z lewej wąską kolumnę na ikonę (kolumna drzewa #0)."""
    frame, tree = scrolled(parent, ttk.Treeview, columns=[c[0] for c in columns],
                           show="tree headings" if icon else "headings", height=height,
                           selectmode="browse")
    if icon:
        tree.heading("#0", text="")
        tree.column("#0", width=44, minwidth=44, stretch=False, anchor="center")
    for cid, head, width, anchor in columns:
        tree.heading(cid, text=head)
        tree.column(cid, width=width, anchor=anchor, stretch=anchor == "w")
    make_sortable(tree, [c[0] for c in columns])
    for k, color in STATUS_COLORS.items():
        if color:
            tree.tag_configure(k, foreground=color)
    tree.tag_configure("hit", background="#fff3b0")
    return frame, tree


class App(tk.Tk):
    def __init__(self, path=None):
        super().__init__()
        self.title(APP_TITLE)
        # ikona okna: obok skryptu albo wewnątrz pliku .exe (PyInstaller rozpakowuje do _MEIPASS)
        ico = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent)) / "kb_icon.ico"
        if ico.is_file():
            try:
                self.iconbitmap(default=str(ico))
            except tk.TclError:
                pass
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        self.geometry(f"{min(1500, sw - 60)}x{min(900, sh - 100)}+20+20")
        self.minsize(980, 600)
        self.cfg = self._load_cfg()
        self.world = None
        self.game = None
        self._game_cache = {}
        self.save_path = None
        self._dup_maps = set()
        self._photos = {}
        self.cur_map = None          # mapa pokazana w podglądzie
        self.hl = None               # wyróżnienie: {"title", "keys", "maps"}
        self.visible_by_map, self.map_iid, self.iid_map = {}, {}, {}
        self.map_combo_ids = []
        self._style()
        self._build()
        start = path or self.cfg.get("last_save")
        if start and Path(start).is_file():
            self.after(50, lambda: self.open_save(start))
        else:
            self.after(50, self._hello)

    # --- konfiguracja i wygląd ---------------------------------------------

    def _load_cfg(self):
        try:
            return json.loads(CONFIG.read_text("utf-8"))
        except (OSError, ValueError):
            return {}

    def _save_cfg(self):
        try:
            CONFIG.write_text(json.dumps(self.cfg, ensure_ascii=False, indent=2), "utf-8")
        except OSError:
            pass

    def _style(self):
        st = ttk.Style(self)
        if "vista" in st.theme_names():
            st.theme_use("vista")
        base = ("Segoe UI", 10)
        st.configure(".", font=base)
        st.configure("Treeview", rowheight=30, font=base)
        st.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))
        st.configure("Title.TLabel", font=("Segoe UI", 13, "bold"))
        st.configure("Sub.TLabel", foreground="#555")
        st.configure("Head.TLabel", font=("Segoe UI", 11, "bold"))
        self._style_tabs(st)
        self.option_add("*TCombobox*Listbox.font", base)

    def _style_tabs(self, st):
        # zakładki motywu vista są rysowane natywnie i ignorują kolory,
        # więc element karty pożyczamy z motywu clam (płaski, da się kolorować)
        try:
            st.element_create("KB.Notebook.tab", "from", "clam")
        except tk.TclError:
            return
        st.layout("TNotebook.Tab", [("KB.Notebook.tab", {"sticky": "nswe", "children": [
            ("Notebook.padding", {"side": "top", "sticky": "nswe", "children": [
                ("Notebook.label", {"side": "top", "sticky": ""})]})]})])
        st.configure("TNotebook", tabmargins=(0, 4, 0, 0))
        st.configure("TNotebook.Tab", font=("Segoe UI", 10, "bold"), padding=(18, 8),
                     background="#dfe5ec", foreground="#3a4654",
                     bordercolor="#b4bfcc", lightcolor="#dfe5ec", darkcolor="#dfe5ec")
        sel, hot = "#0b5cad", "#c9d7e8"
        st.map("TNotebook.Tab",
               background=[("selected", sel), ("active", hot)],
               lightcolor=[("selected", sel), ("active", hot)],
               darkcolor=[("selected", sel), ("active", hot)],
               bordercolor=[("selected", sel)],
               foreground=[("selected", "white"), ("active", "#0b3f78")],
               expand=[("selected", (0, 2, 0, 0))])

    # --- układ okna -----------------------------------------------------------

    def _build(self):
        bar = ttk.Frame(self, padding=(10, 8, 10, 4))
        bar.pack(fill="x")
        ttk.Button(bar, text="Otwórz zapis…", command=self.ask_open).pack(side="left")
        ttk.Button(bar, text="Wczytaj ponownie", command=self.reload).pack(side="left", padx=6)
        ttk.Button(bar, text="Eksport do Excela…", command=self.export).pack(side="left")
        ttk.Button(bar, text="Folder gry…", command=self.ask_game_dir).pack(side="right")
        self.game_lbl = ttk.Label(bar, style="Sub.TLabel")
        self.game_lbl.pack(side="right", padx=8)

        head = ttk.Frame(self, padding=(10, 2, 10, 6))
        head.pack(fill="x")
        self.title_lbl = ttk.Label(head, text="Nie wczytano zapisu", style="Title.TLabel")
        self.title_lbl.pack(anchor="w")
        self.sub_lbl = ttk.Label(head, text="", style="Sub.TLabel")
        self.sub_lbl.pack(anchor="w")

        self.nb = ttk.Notebook(self)
        self.nb.pack(fill="both", expand=True, padx=10, pady=(0, 4))
        self._tab_places()
        self._tab_items()
        self._tab_spells()
        self._tab_units()
        self._tab_hero()

        self.status = ttk.Label(self, padding=(10, 2, 10, 6), style="Sub.TLabel")
        self.status.pack(fill="x")

    def _filter_bar(self, parent, on_change):
        f = ttk.Frame(parent, padding=(0, 6, 0, 6))
        ttk.Label(f, text="Szukaj:").pack(side="left")
        var = tk.StringVar()
        e = ttk.Entry(f, textvariable=var, width=30)
        e.pack(side="left", padx=(4, 10))
        var.trace_add("write", lambda *_: on_change())
        ttk.Button(f, text="×", width=3, command=lambda: var.set("")).pack(side="left", padx=(0, 12))
        return f, var

    # zakładka 1: mapy i budynki
    def _tab_places(self):
        tab = ttk.Frame(self.nb, padding=6)
        self.nb.add(tab, text="Mapy i budynki")
        bar, self.pl_q = self._filter_bar(tab, self.fill_places)
        bar.pack(fill="x")
        ttk.Label(bar, text="Mapy:").pack(side="left")
        self.pl_maps = ttk.Combobox(bar, state="readonly", width=16,
                                    values=["wszystkie", "odwiedzone", "nieodwiedzone"])
        self.pl_maps.current(0)
        self.pl_maps.bind("<<ComboboxSelected>>", lambda e: self.fill_places())
        self.pl_maps.pack(side="left", padx=(4, 12))
        self.pl_kinds = {}
        for kind, label in KINDS:
            v = tk.BooleanVar(value=kind != RES)
            ttk.Checkbutton(bar, text=label, variable=v, command=self.fill_places).pack(side="left", padx=4)
            self.pl_kinds[kind] = v

        # układ: po lewej lista miejsc nad zawartością wybranego miejsca, po prawej duża mapa
        pane = ttk.PanedWindow(tab, orient="horizontal")
        pane.pack(fill="both", expand=True)
        left = ttk.PanedWindow(pane, orient="vertical")
        fr_tree, self.pl_tree = scrolled(left, ttk.Treeview, columns=("info",), show="tree headings",
                                         selectmode="browse", height=12)
        self.pl_tree.heading("#0", text="Mapa / miejsce")
        self.pl_tree.heading("info", text="Zawartość")
        self.pl_tree.column("#0", width=300)
        self.pl_tree.column("info", width=210)
        self.pl_tree.tag_configure("map", font=("Segoe UI", 10, "bold"))
        self.pl_tree.tag_configure("unvisited", foreground="#6b4e00")
        self.pl_tree.tag_configure("group", font=("Segoe UI", 10, "bold"), foreground="#0b5cad")
        self.pl_tree.bind("<<TreeviewSelect>>", lambda e: self.show_place())
        left.add(fr_tree, weight=3)

        bottom = ttk.Frame(left, padding=(0, 8, 0, 0))
        self.pl_head = ttk.Label(bottom, text="Wybierz miejsce z listy", style="Head.TLabel")
        self.pl_head.pack(anchor="w")
        self.pl_sub = ttk.Label(bottom, text="", style="Sub.TLabel", wraplength=560, justify="left")
        self.pl_sub.pack(anchor="w", pady=(0, 6))
        fr, self.pl_items = table(bottom, [
            ("kind", "Rodzaj", 95, "w"), ("name", "Nazwa", 190, "w"), ("count", "Ilość", 60, "e"),
            ("type", "Typ / szkoła", 115, "w"), ("level", "Poziom", 60, "center"),
            ("role", "Jak zdobyć", 120, "w"), ("note", "Uwagi", 120, "w")], height=7, icon=True)
        fr.pack(fill="both", expand=True)
        self.pl_items.bind("<Double-1>", lambda e: self._goto_entry(self.pl_items))
        bottom.bind("<Configure>", lambda e: self.pl_sub.config(wraplength=max(200, e.width - 10)))
        left.add(bottom, weight=2)
        pane.add(left, weight=2)

        right = ttk.Frame(pane, padding=(8, 0, 0, 0))
        mbar = ttk.Frame(right)
        mbar.pack(fill="x", pady=(0, 4))
        ttk.Label(mbar, text="Mapa:").pack(side="left")
        self.map_combo = ttk.Combobox(mbar, state="readonly", width=38)
        self.map_combo.bind("<<ComboboxSelected>>", lambda e: self._map_combo_pick())
        self.map_combo.pack(side="left", padx=(4, 10))
        self.mapview = MapView(right, on_pick=self._map_pick, label_fn=self._marker_label)
        self.mapview.pack(fill="both", expand=True)
        # wiersz z wyróżnieniem ("Pokaż na mapie") - widoczny tylko, gdy coś jest wyróżnione
        self.map_hl_bar = ttk.Frame(right)
        self.map_hl_lbl = ttk.Label(self.map_hl_bar, text="", foreground="#8a6d00", font=("Segoe UI", 10, "bold"))
        self.map_hl_lbl.pack(side="left")
        ttk.Button(self.map_hl_bar, text="× Wyczyść wyróżnienie",
                   command=self.clear_highlight).pack(side="left", padx=10)
        pane.add(right, weight=3)

        def init_sashes(event):
            # przy pierwszym wyświetleniu: mapa (kwadratowa) dostaje szerokość równą swojej wysokości,
            # lista resztę, ale nie mniej niż 30% okna
            if event.width > 100 and not getattr(self, "_sashes_set", False):
                self._sashes_set = True
                map_w = event.height - mbar.winfo_reqheight() + 16
                pane.sashpos(0, max(int(event.width * 0.30), event.width - map_w))
                left.after_idle(lambda: left.sashpos(0, int(left.winfo_height() * 0.55)))
        pane.bind("<Configure>", init_sashes, add="+")
        self.place_by_iid = {}
        self.iid_by_place = {}

    # zakładka 2: przedmioty
    def _tab_items(self):
        tab = ttk.Frame(self.nb, padding=6)
        self.nb.add(tab, text="Przedmioty")
        bar, self.it_q = self._filter_bar(tab, self.fill_items)
        bar.pack(fill="x")
        ttk.Label(bar, text="Typ:").pack(side="left")
        self.it_type = ttk.Combobox(bar, state="readonly", width=16)
        self.it_type.bind("<<ComboboxSelected>>", lambda e: self.fill_items())
        self.it_type.pack(side="left", padx=(4, 12))
        ttk.Label(bar, text="Status:").pack(side="left")
        self.it_status = ttk.Combobox(bar, state="readonly", width=24,
                                      values=[l for _, l in STATUS_LABELS])
        self.it_status.current(0)
        self.it_status.bind("<<ComboboxSelected>>", lambda e: self.fill_items())
        self.it_status.pack(side="left", padx=(4, 12))
        bar = ttk.Frame(tab, padding=(0, 0, 0, 6))
        bar.pack(fill="x")
        ttk.Label(bar, text="Premia:").pack(side="left")
        self.it_bonus = ttk.Combobox(bar, state="readonly", width=34, values=["wszystkie"])
        self.it_bonus.current(0)
        self.it_bonus.bind("<<ComboboxSelected>>", lambda e: self.fill_items())
        self.it_bonus.pack(side="left", padx=(4, 12))
        self.it_desc_q = tk.BooleanVar(value=False)
        ttk.Checkbutton(bar, text="Szukaj też w opisie", variable=self.it_desc_q,
                        command=self.fill_items).pack(side="left", padx=(0, 10))
        self.it_all = tk.BooleanVar(value=False)
        ttk.Checkbutton(bar, text="Pokaż medale i premie ukryte", variable=self.it_all,
                        command=self.fill_items).pack(side="left")

        pane = ttk.PanedWindow(tab, orient="vertical")
        pane.pack(fill="both", expand=True)
        fr, self.it_tree = table(pane, [
            ("name", "Nazwa", 230, "w"), ("type", "Typ", 110, "w"), ("bonus", "Premie", 260, "w"),
            ("level", "Poziom", 60, "center"),
            ("race", "Rasa", 90, "w"), ("price", "Cena", 70, "e"), ("max", "Limit w grze", 105, "e"),
            ("box", "Wylosowano", 100, "e"), ("where", "Miejsc", 65, "e"), ("status", "Czy pojawi się w grze?", 380, "w")],
            height=8, icon=True)
        self.it_tree.bind("<<TreeviewSelect>>", lambda e: self.show_item())
        pane.add(fr, weight=1)
        self.it_desc, self.it_where = self._detail_pane(pane, lambda: self._map_from(self.it_tree, [ITEM]))

    def _detail_pane(self, pane, on_map=None):
        """Dolny panel: obrazek i opis po lewej, lista miejsc występowania po prawej."""
        low = ttk.PanedWindow(pane, orient="horizontal")
        fr, txt = scrolled(low, tk.Text, width=58, height=17, wrap="word", font=("Segoe UI", 10),
                           relief="flat", background="#f7f7f7", padx=10, pady=8)
        txt.tag_configure("h", font=("Segoe UI", 13, "bold"), spacing1=4, spacing3=2)
        txt.tag_configure("h2", font=("Segoe UI", 10, "bold"), spacing1=8, spacing3=2)
        txt.tag_configure("s", foreground="#555")
        txt.tag_configure("i", font=("Segoe UI", 10, "italic"), foreground="#444")
        txt.tag_configure("b", font=("Segoe UI", 10, "bold"))
        txt.tag_configure("li", lmargin1=12, lmargin2=24)
        for k, color in STATUS_COLORS.items():
            if color:
                txt.tag_configure(k, foreground=color, font=("Segoe UI", 10, "bold"))
        txt.config(state="disabled")
        low.add(fr, weight=2)
        side = ttk.Frame(low)
        if on_map:
            bar = ttk.Frame(side)
            bar.pack(fill="x", pady=(0, 4))
            ttk.Button(bar, text="Pokaż na mapie", command=on_map).pack(side="left")
            ttk.Label(bar, text="dwuklik na wierszu – przejście do miejsca na mapie",
                      style="Sub.TLabel").pack(side="left", padx=8)
        fr2, where = table(side, self._where_columns(), height=8)
        fr2.pack(fill="both", expand=True)
        where.bind("<Double-1>", lambda e: self._goto_where(where))
        low.add(side, weight=3)
        pane.add(low, weight=1)
        return txt, where

    def _where_columns(self):
        return [("map", "Mapa", 190, "w"), ("place", "Miejsce", 260, "w"), ("role", "Jak zdobyć", 130, "w"),
                ("count", "Ilość", 60, "e"), ("state", "Stan mapy", 230, "w")]

    # zakładka 3: czary
    def _tab_spells(self):
        tab = ttk.Frame(self.nb, padding=6)
        self.nb.add(tab, text="Czary")
        bar, self.sp_q = self._filter_bar(tab, self.fill_spells)
        bar.pack(fill="x")
        ttk.Label(bar, text="Szkoła:").pack(side="left")
        self.sp_school = ttk.Combobox(bar, state="readonly", width=16,
                                      values=["wszystkie"] + list(SZKOLY.values()))
        self.sp_school.current(0)
        self.sp_school.bind("<<ComboboxSelected>>", lambda e: self.fill_spells())
        self.sp_school.pack(side="left", padx=(4, 12))
        pane = ttk.PanedWindow(tab, orient="vertical")
        pane.pack(fill="both", expand=True)
        fr, self.sp_tree = table(pane, [
            ("name", "Czar", 220, "w"), ("school", "Szkoła", 130, "w"), ("level", "Poziom", 70, "center"),
            ("price", "Cena", 75, "e"), ("book", "W księdze", 95, "center"), ("scrolls", "Twoje zwoje", 110, "e"),
            ("where", "Miejsc", 70, "e"), ("status", "Status", 300, "w")], height=8, icon=True)
        self.sp_tree.bind("<<TreeviewSelect>>", lambda e: self.show_spell())
        pane.add(fr, weight=1)
        self.sp_desc, self.sp_where = self._detail_pane(pane, lambda: self._map_from(self.sp_tree, [SPELL]))

    # zakładka 4: jednostki
    def _tab_units(self):
        tab = ttk.Frame(self.nb, padding=6)
        self.nb.add(tab, text="Jednostki")
        bar, self.un_q = self._filter_bar(tab, self.fill_units)
        bar.pack(fill="x")
        self.un_enemy = tk.BooleanVar(value=False)
        ttk.Checkbutton(bar, text="Także jednostki tylko w armiach wroga", variable=self.un_enemy,
                        command=self.fill_units).pack(side="left")
        pane = ttk.PanedWindow(tab, orient="vertical")
        pane.pack(fill="both", expand=True)
        fr, self.un_tree = table(pane, [
            ("name", "Jednostka", 220, "w"), ("shops", "Miejsc werbunku", 150, "e"),
            ("total", "Sztuk do zwerbowania", 190, "e"), ("join", "Dołączą do armii", 160, "e"),
            ("enemy", "Armii wroga", 120, "e"), ("mine", "W twojej armii", 140, "e")], height=8, icon=True)
        self.un_tree.bind("<<TreeviewSelect>>", lambda e: self.show_unit())
        pane.add(fr, weight=1)
        self.un_desc, self.un_where = self._detail_pane(pane, lambda: self._map_from(self.un_tree, [UNIT, ARMY]))

    # zakładka 5: bohater
    def _tab_hero(self):
        tab = ttk.Frame(self.nb, padding=6)
        self.nb.add(tab, text="Bohater")
        fr, self.hero_txt = scrolled(tab, tk.Text, wrap="word", font=("Consolas", 10), relief="flat",
                                     padx=10, pady=8)
        fr.pack(fill="both", expand=True)

    # --- wczytywanie ---------------------------------------------------------

    def _hello(self):
        self.status.config(text="Otwórz plik zapisu (.sav) przyciskiem „Otwórz zapis…”.")
        self._update_game_label()

    def ask_open(self):
        initial = self.cfg.get("last_save")
        initdir = str(Path(initial).parent) if initial else next((str(d) for d in save_dirs()), None)
        path = filedialog.askopenfilename(
            title="Wybierz zapis King's Bounty", initialdir=initdir,
            filetypes=[("Zapisy King's Bounty", "*.sav"), ("Wszystkie pliki", "*.*")])
        if path:
            self.open_save(path)

    def reload(self):
        if self.save_path:
            self.open_save(self.save_path)

    def ask_game_dir(self):
        d = filedialog.askdirectory(title="Wskaż folder gry (z podfolderem 'sessions')",
                                    initialdir=self.cfg.get("game_dir") or None)
        if not d:
            return
        if not is_game_dir(d):
            messagebox.showwarning(APP_TITLE, "W tym folderze nie ma podfolderu „sessions”.\n"
                                              "Wskaż główny folder gry (tam, gdzie jest KB.exe).")
            return
        self.cfg["game_dir"] = d
        self._save_cfg()
        self._game_cache.clear()
        self.reload()
        self._update_game_label()

    def _game_dir(self):
        d = self.cfg.get("game_dir")
        if d and is_game_dir(d):
            return d
        found = find_game_dirs()
        if found:
            self.cfg["game_dir"] = str(found[0])
            self._save_cfg()
            return str(found[0])
        return None

    def _update_game_label(self):
        d = self._game_dir()
        self.game_lbl.config(text=f"Folder gry: {d}" if d else "Nie znaleziono folderu gry – nazwy będą identyfikatorami")

    def open_save(self, path):
        self.config(cursor="watch")
        self.update_idletasks()
        try:
            info, data = load_save(path)
            session = info.get("session") or "addon"
            key = (self._game_dir(), session)
            if key not in self._game_cache:
                self._game_cache[key] = GameData(key[0], session)
            if self.game is not self._game_cache[key]:
                self._photos.clear()
            self.game = self._game_cache[key]
            self.world = World(info, data, self.game)
        except Exception as e:  # noqa: BLE001 - pokazujemy użytkownikowi każdy błąd odczytu
            self.config(cursor="")
            messagebox.showerror(APP_TITLE, f"Nie udało się odczytać zapisu:\n{path}\n\n{e}")
            return
        self.config(cursor="")
        self.save_path = path
        self.cfg["last_save"] = str(path)
        self._save_cfg()
        self._update_game_label()
        self._after_load()

    def _after_load(self):
        w, g = self.world, self.game
        names = [g.map_name(m) for m, _ in w.maps()]
        self._dup_maps = {n for n in names if names.count(n) > 1}
        h, z = w.hero["bohater"], w.hero["zapis"]
        self.title(f"{APP_TITLE} – {Path(self.save_path).name}")
        self.title_lbl.config(text=f"{h['imie']} – {h['klasa'].split('(')[-1].rstrip(')')}, poziom {h['poziom']}")
        mtime = Path(self.save_path).stat().st_mtime
        import datetime
        when = datetime.datetime.fromtimestamp(mtime).strftime("%Y-%m-%d %H:%M")
        self.sub_lbl.config(text=f"Dzień {z['dzien_gry']}  •  lokacja: {g.map_name(z['lokacja'] or '')}  •  "
                                 f"złoto {h['zloto']}  •  przywództwo {h['przywodztwo']}  •  "
                                 f"plik: {self.save_path} (zapisany {when})")
        self.hl = None
        self.map_hl_lbl.config(text="")
        self.map_hl_bar.pack_forget()
        self.cur_map = None
        self._start_map = z["lokacja"] if any(m == z["lokacja"] for m, _ in w.maps()) else None
        types = sorted({self.item_type(i) for i in self._item_ids()})
        self.it_type.config(values=["wszystkie"] + types)
        self.it_type.current(0)
        labels = {}
        for it in g.items.values():
            for bn in it.get("bonuses", []):
                if bonus_good(bn):
                    labels.setdefault(bn["label"], GRUPY_PREMII.index(bn["group"]))
        self.it_bonus.config(values=["wszystkie", "dowolna premia dla bohatera", "dowolna premia dla wojsk"]
                             + sorted(labels, key=lambda l: (labels[l], l)))
        self.it_bonus.current(0)
        self.fill_places()
        self.fill_items()
        self.fill_spells()
        self.fill_units()
        self.fill_hero()
        maps = w.maps()
        nv = sum(1 for m, _ in maps if m in w.visited)
        msg = (f"Mapy: {len(maps)} (odwiedzone {nv}, pozostałe {len(maps) - nv} – ich zawartość gra już wylosowała)  •  "
               f"miejsc z zawartością: {len(w.places)}")
        if not g.ok:
            msg += "  •  " + (g.error or "brak plików gry: nazwy są identyfikatorami, brak typów i poziomów")
        self.status.config(text=msg)

    # --- obrazki ---------------------------------------------------------------

    def photo(self, kind, iid, box, big=False):
        """Obrazek z gry przeskalowany do ramki box=(szer., wys.) albo None."""
        key = (kind, iid, box, big)
        if key not in self._photos:
            img = None
            if ImageTk and self.game:
                g = self.game
                if kind == ITEM:
                    src = g.item_image(iid)
                elif kind == SPELL:
                    src = g.spell_image(iid, scroll=not big) or g.spell_image(iid, scroll=big)
                elif kind in (UNIT, ARMY):
                    src = g.unit_image(iid)
                else:
                    src = g.image(RES_ICONS.get(iid))
                if src:
                    scale = min(box[0] / src.width, box[1] / src.height)
                    size = (max(1, round(src.width * scale)), max(1, round(src.height * scale)))
                    img = ImageTk.PhotoImage(src.resize(size, Image.LANCZOS))
            self._photos[key] = img
        return self._photos[key]

    def icon_kw(self, kind, iid):
        ph = self.photo(kind, iid, (40, 26))
        return {"image": ph} if ph else {}

    def _text_begin(self, txt):
        txt.config(state="normal")
        txt.delete("1.0", "end")

    def _text_image(self, txt, *photos):
        shown = [p for p in photos if p]
        for ph in shown:
            txt.image_create("end", image=ph, padx=2, pady=2)
        if shown:
            txt.insert("end", "\n")

    # --- nazwy ---------------------------------------------------------------

    def ename(self, kind, iid):
        g = self.game
        if kind == ITEM:
            return g.item_name(iid)
        if kind == SPELL:
            return g.spell_name(iid)
        if kind in (UNIT, ARMY):
            return g.unit_name(iid)
        return g.resource_name(iid)

    def item_type(self, iid):
        it = self.game.items.get(iid)
        if not it:
            return "Nieznany"
        if it["slot"]:
            return SLOTY[it["slot"]]
        bits = set(it["bits"])
        if "wife" in bits:
            return "Towarzysz"
        if "container" in bits:
            return "Jajo / mikstura"
        if "quest" in bits:
            return "Zadaniowy"
        if "usable" in bits:
            return "Do użycia"
        return "Inne"

    def entry_type(self, e):
        if e.kind == ITEM:
            it = self.game.items.get(e.id) or {}
            return self.item_type(e.id), it.get("level", "")
        if e.kind == SPELL:
            sp = self.game.spells.get(e.id) or {}
            return SZKOLY.get(sp.get("school"), ""), sp.get("level", "")
        return "", ""

    def place_label(self, p):
        if p.name:
            return f"{p.type}: {p.name}"
        if p.atom and p.type in BUILDING_TYPES:
            return f"{p.type}: {pretty_atom(p.atom)}"
        ents = [e for e in p.entries if e.kind != RES] or p.entries
        first = self.ename(ents[0].kind, ents[0].id)
        more = f" (+{len(ents) - 1})" if len(ents) > 1 else ""
        return f"{p.type}: {first}{more}"

    def map_state(self, p):
        if p.source == "zadanie":
            return f"zadanie – {p.atom}" if p.atom else "dialog z postacią"
        if p.source == "bohater":
            return "u ciebie"
        return p.source

    def map_label(self, p):
        if p.source == "zadanie":
            return "Zadania i dialogi"
        if p.source == "bohater":
            return "Bohater"
        return self.map_title(p.map_id)

    def map_title(self, map_id):
        """Nazwa mapy z gry; kilka map ma tę samą nazwę, wtedy dopisany jest identyfikator."""
        name = self.game.map_name(map_id)
        return f"{name} ({map_id})" if name in self._dup_maps else name

    # --- zakładka 1 ------------------------------------------------------------

    def fill_places(self):
        if not self.world:
            return
        t = self.pl_tree
        t.delete(*t.get_children())
        self.place_by_iid, self.iid_by_place = {}, {}
        self.visible_by_map, self.map_iid, self.iid_map = defaultdict(list), {}, {}
        q = self.pl_q.get().strip().casefold()
        kinds = {k for k, v in self.pl_kinds.items() if v.get()}
        mode = self.pl_maps.get()
        w = self.world
        groups = [(m, self.map_title(m), ps) for m, ps in w.maps()]
        groups.append(("__quests__", "Zadania i dialogi (nagrody wylosowane z góry)",
                       [p for p in w.places if p.source == "zadanie"]))
        groups.append(("__hero__", "Bohater", [p for p in w.places if p.source == "bohater"]))
        for map_id, map_name, places in groups:
            special = map_id.startswith("__")
            visited = map_id in w.visited
            if not special and ((mode == "odwiedzone" and not visited) or (mode == "nieodwiedzone" and visited)):
                continue
            if special and mode != "wszystkie":
                continue
            rows = []
            for p in places:
                ents = [e for e in p.entries if e.kind in kinds]
                if not ents:
                    continue
                label = self.place_label(p)
                if q and q not in label.casefold() and not any(
                        q in self.ename(e.kind, e.id).casefold() or q in e.id.casefold() for e in ents):
                    continue
                rows.append((p, label, ents))
            if not rows:
                continue
            state = f"{len(rows)} {miejsc(len(rows))}" + (
                "" if special else (", odwiedzona" if visited else ", nieodwiedzona"))
            tags = ("group",) if special else ("map",) if visited else ("map", "unvisited")
            mnode = t.insert("", "end", text=map_name, values=(state,), open=bool(q), tags=tags)
            if not special:
                self.map_iid[map_id], self.iid_map[mnode] = mnode, map_id
            rows.sort(key=lambda r: (TYPE_ORDER.index(r[0].type) if r[0].type in TYPE_ORDER else 99,
                                     r[1].casefold()))
            for p, label, ents in rows:
                counts = {}
                for e in ents:
                    counts[e.kind] = counts.get(e.kind, 0) + 1
                info = ", ".join(f"{n} {SHORT[k]}" for k, n in counts.items())
                iid = t.insert(mnode, "end", text=label, values=(info,))
                self.place_by_iid[iid] = p
                self.iid_by_place[p.key] = iid
                if p.map_id:
                    self.visible_by_map[p.map_id].append(p)
        self._update_map_combo()
        self.show_place()

    def show_place(self):
        sel = self.pl_tree.selection()
        p = self.place_by_iid.get(sel[0]) if sel else None
        tv = self.pl_items
        tv.delete(*tv.get_children())
        if not p:
            map_id = self.iid_map.get(sel[0]) if sel else None
            if map_id:
                n = len(self.visible_by_map.get(map_id, []))
                state = "odwiedzona" if map_id in self.world.visited else "nieodwiedzona – zawartość już wylosowana"
                self.pl_head.config(text=self.map_title(map_id))
                self.pl_sub.config(text=f"{state}  •  {n} {miejsc(n)} na mapie  •  "
                                        "kliknij punkt na mapie albo miejsce na liście")
            else:
                self.pl_head.config(text="Wybierz miejsce z listy")
                self.pl_sub.config(text="Mapy odwiedzone pokazują stan bieżący (to, co jeszcze zostało). "
                                        "Mapy nieodwiedzone pokazują to, co gra już wylosowała na początku rozgrywki.")
            self.show_map(map_id or self.cur_map or getattr(self, "_start_map", None))
            return
        self.show_map(p.map_id, selected=p.key)
        self.pl_head.config(text=self.place_label(p))
        sub = [self.map_label(p), self.map_state(p)]
        if p.atom and p.source not in ("zadanie",):
            sub.append(f"obiekt gry: {p.atom}")
        sub.append(f"id: {p.key}")
        self.pl_sub.config(text="  •  ".join(s for s in sub if s))
        kinds = {k for k, v in self.pl_kinds.items() if v.get()}
        q = self.pl_q.get().strip().casefold()
        for i, e in enumerate(p.entries):
            if e.kind not in kinds:
                continue
            name = self.ename(e.kind, e.id)
            typ, lvl = self.entry_type(e)
            tags = ["hit"] if q and (q in name.casefold() or q in e.id.casefold()) else []
            if e.kind == ITEM:
                tags.append(self.world.item_status(e.id)[0])
            tv.insert("", "end", iid=f"e{i}", values=(e.kind, name, e.count, typ, lvl,
                                                      ROLA.get(e.role, e.role), e.note), tags=tags,
                      **self.icon_kw(e.kind, e.id))

    # --- podgląd mapy ------------------------------------------------------------

    def show_map(self, map_id, selected=None):
        mv, g = self.mapview, self.game
        if not map_id or not g:
            mv.show(None, [], message="To miejsce nie ma pozycji na mapie (nagroda za zadanie, rozmowa albo bohater)."
                    if selected or self.pl_tree.selection() else "Wybierz mapę lub miejsce z listy.")
            return
        places = list(self.visible_by_map.get(map_id, []))
        if self.hl:
            have = {p.key for p in places}
            places += [p for p in self.world.places
                       if p.map_id == map_id and p.key in self.hl["keys"] and p.key not in have]
        markers = []
        for p in places:
            uv = g.map_uv(map_id, p.pos)
            if uv:
                markers.append((p, uv[0], uv[1]))
        img = g.radar_image(map_id)
        if img:
            msg = ""
        elif ImageTk is None:
            msg = "Podgląd mapy wymaga biblioteki Pillow (pip install pillow)."
        else:
            msg = "Brak minimapy tej lokacji w plikach gry (albo nie wskazano folderu gry)."
        keep = map_id == self.cur_map
        self.cur_map = map_id
        mv.show(img, markers, highlight=self.hl["keys"] if self.hl else (), selected=selected,
                keep_view=keep, message=msg)
        if selected and keep:
            mv.center_on(selected)
        if map_id in self.map_combo_ids:
            self.map_combo.current(self.map_combo_ids.index(map_id))

    def _update_map_combo(self):
        if not self.world:
            return
        ids = [m for m, _ in self.world.maps()]
        counts = self.hl["maps"] if self.hl else {}
        ids.sort(key=lambda m: -counts.get(m, 0))
        self.map_combo_ids = ids
        self.map_combo.config(values=[self.map_title(m) + (f"   ★ {counts[m]}" if counts.get(m) else "")
                                      for m in ids])
        if self.cur_map in ids:
            self.map_combo.current(ids.index(self.cur_map))

    def _map_combo_pick(self):
        i = self.map_combo.current()
        if i < 0:
            return
        map_id = self.map_combo_ids[i]
        node = self.map_iid.get(map_id)
        if node:
            self.pl_tree.selection_set(node)
            self.pl_tree.see(node)
        else:
            self.show_map(map_id)

    def _map_pick(self, p):
        self.goto_place(p.key)

    def _marker_label(self, p):
        names = []
        for e in p.entries:
            if e.kind == RES:
                continue
            n = self.ename(e.kind, e.id)
            if n not in names:
                names.append(n)
        extra = ", ".join(names[:4]) + (f" … (+{len(names) - 4})" if len(names) > 4 else "")
        return self.place_label(p) + (f"\n{extra}" if extra else "")

    def _map_from(self, tree, kinds):
        sel = tree.selection()
        if sel:
            self.show_on_map(kinds, sel[0], tree.set(sel[0], "name"))

    def show_on_map(self, kinds, iid, title):
        """Wyróżnia na mapach wszystkie miejsca, gdzie występuje przedmiot / czar / jednostka."""
        by_map = defaultdict(list)
        for k in kinds:
            for p, _ in self.world.by_id.get((k, iid), []):
                if p.map_id and p.pos and p.key not in by_map[p.map_id]:
                    by_map[p.map_id].append(p.key)
        if not by_map:
            messagebox.showinfo(APP_TITLE, f"{title}: nie występuje w żadnym miejscu na mapach.\n"
                                           "Może być tylko nagrodą za zadanie, u bohatera albo wcale.")
            return
        n = sum(len(v) for v in by_map.values())
        self.hl = {"title": title, "keys": {k for v in by_map.values() for k in v},
                   "maps": {m: len(v) for m, v in by_map.items()}}
        self.map_hl_lbl.config(text=f"★ {title}: {n} {miejsc(n)} na "
                                    f"{len(by_map)} {'mapie' if len(by_map) == 1 else 'mapach'}")
        self.map_hl_bar.pack(fill="x", pady=(0, 4), before=self.mapview)
        self._update_map_combo()
        best = max(by_map, key=lambda m: len(by_map[m]))
        self.cur_map = None
        self.goto_place(by_map[best][0])

    def clear_highlight(self):
        self.hl = None
        self.map_hl_lbl.config(text="")
        self.map_hl_bar.pack_forget()
        self._update_map_combo()
        sel = self.pl_tree.selection()
        p = self.place_by_iid.get(sel[0]) if sel else None
        self.show_map(self.cur_map, selected=p.key if p else None)

    def goto_place(self, key):
        iid = self.iid_by_place.get(key)
        if not iid:
            # miejsce ukryte filtrami - pokaż wszystko
            self.pl_q.set("")
            self.pl_maps.current(0)
            for v in self.pl_kinds.values():
                v.set(True)
            self.fill_places()
            iid = self.iid_by_place.get(key)
        if iid:
            self.nb.select(0)
            self.pl_tree.item(self.pl_tree.parent(iid), open=True)
            self.pl_tree.selection_set(iid)
            self.pl_tree.see(iid)

    def _goto_entry(self, tv):
        """Podwójne kliknięcie przedmiotu w zawartości miejsca: przejście do jego karty."""
        sel = tv.selection()
        if not sel:
            return
        kind, name = tv.set(sel[0], "kind"), tv.set(sel[0], "name")
        target = {ITEM: (1, self.it_tree), SPELL: (2, self.sp_tree), UNIT: (3, self.un_tree),
                  ARMY: (3, self.un_tree)}.get(kind)
        if not target:
            return
        tab, tree = target
        if tab == 1:
            self.it_q.set(""); self.it_type.current(0); self.it_status.current(0)
            self.it_all.set(True); self.fill_items()
        elif tab == 2:
            self.sp_q.set(""); self.sp_school.current(0); self.fill_spells()
        else:
            self.un_q.set(""); self.un_enemy.set(True); self.fill_units()
        for k in tree.get_children():
            if tree.set(k, "name") == name:
                self.nb.select(tab)
                tree.selection_set(k)
                tree.see(k)
                break

    # --- zakładka 2 ------------------------------------------------------------

    def _item_ids(self):
        ids = set(self.game.items)
        ids |= {i for (k, i) in self.world.by_id if k == ITEM}
        return ids

    def fill_items(self):
        if not self.world:
            return
        t = self.it_tree
        t.delete(*t.get_children())
        self._item_score = {}
        q = self.it_q.get().strip().casefold()
        typ = self.it_type.get()
        st_label = self.it_status.get()
        st_key = next((k for k, l in STATUS_LABELS if l == st_label), "")
        show_all = self.it_all.get()
        want = self.it_bonus.get()
        in_desc = self.it_desc_q.get()
        n = 0
        for iid in self._item_ids():
            it = self.game.items.get(iid) or {}
            if not show_all and it.get("slot") in HIDDEN_SLOTS:
                continue
            bonuses = it.get("bonuses", [])
            score = None
            if want not in ("", "wszystkie"):
                good = [bn for bn in bonuses if bonus_good(bn)]
                if want == "dowolna premia dla bohatera":
                    hits = [bn for bn in good if bn["group"] == "bohater"]
                elif want == "dowolna premia dla wojsk":
                    hits = [bn for bn in good if bn["group"] != "bohater" and bn["group"] != "zdolność"]
                else:
                    hits = [bn for bn in good if bn["label"] == want]
                if not hits:
                    continue
                score = max(abs(bn["value"]) for bn in hits)
            name = self.game.item_name(iid)
            ity = self.item_type(iid)
            if typ not in ("", "wszystkie") and ity != typ:
                continue
            key, text = self.world.item_status(iid)
            if st_key and key != st_key:
                continue
            if q and q not in name.casefold() and q not in iid.casefold() and not (
                    in_desc and q in self.game.item_description(iid).casefold()):
                continue
            where = len([x for x in self.world.by_id.get((ITEM, iid), []) if x[0].source != "bohater"])
            btxt = ", ".join(bonus_text(bn).replace(" (niektóre jednostki)", "*") for bn in bonuses)
            self._item_score[iid] = score
            t.insert("", "end", iid=iid, tags=(key,), **self.icon_kw(ITEM, iid), values=(
                name, ity, btxt, it.get("level", ""), RASY.get(it.get("race"), it.get("race", "")),
                it.get("price") or "", it.get("maxcount") if it.get("maxcount") is not None else "",
                self.world.box.get(iid, 0), where or "", text))
            n += 1
        if want not in ("", "wszystkie"):
            # przy filtrze premii najmocniejsze na górze
            rows = sorted(t.get_children(), key=lambda k: (-(self._item_score.get(k) or 0), t.set(k, "name").casefold()))
        else:
            rows = sorted(t.get_children(), key=lambda k: t.set(k, "name").casefold())
        for i, k in enumerate(rows):
            t.move(k, "", i)
        self.nb.tab(1, text=f"Przedmioty ({n})")
        self.show_item()

    def show_item(self):
        sel = self.it_tree.selection()
        txt = self.it_desc
        self._text_begin(txt)
        self.it_where.delete(*self.it_where.get_children())
        if not sel:
            txt.insert("end", "Wybierz przedmiot z listy.\n\n", "h")
            txt.insert("end", "„Wylosowano” to licznik gry: ile egzemplarzy przedmiotu wylosowała do skrzyń, "
                              "sklepów i nagród na wszystkich mapach. Gra losuje to na początku rozgrywki, "
                              "więc przedmiot z zerem i bez miejsca na liście już się nie pojawi "
                              "(chyba że da go zadanie, rozmowa albo ulepszenie innego przedmiotu).", "s")
            txt.config(state="disabled")
            return
        iid = sel[0]
        g, it = self.game, self.game.items.get(sel[0]) or {}
        self._text_image(txt, self.photo(ITEM, iid, (104, 104), big=True))
        txt.insert("end", g.item_name(iid) + "\n", "h")
        facts = [self.item_type(iid)]
        if it.get("level"):
            facts.append(f"poziom {it['level']}")
        if it.get("race"):
            facts.append(RASY.get(it["race"], it["race"]))
        if it.get("price"):
            facts.append(f"cena {it['price']}")
        if it.get("maxcount") is not None:
            facts.append(f"limit w grze: {it['maxcount']}")
        txt.insert("end", ", ".join(facts) + f"\nid: {iid}\n", "s")
        status_key, status_text = self.world.item_status(iid)
        txt.insert("end", "\n" + status_text + "\n", status_key)
        if it.get("bonuses"):
            txt.insert("end", "Premie\n", "h2")
            for bn in it["bonuses"]:
                txt.insert("end", "• " + bonus_text(bn) + "\n", ("li",) if bonus_good(bn) else ("li", "brak"))
        if it.get("upgrade"):
            txt.insert("end", f"Można ulepszyć do: {g.item_name(it['upgrade'])}\n", "s")
        for src in self.world.upgraded_from.get(iid, []):
            txt.insert("end", f"Powstaje z ulepszenia: {g.item_name(src)}\n", "s")
        if it.get("set"):
            txt.insert("end", f"Część kompletu: {it['set']}\n", "s")
        desc = g.item_description(iid)
        if desc:
            hint, _, info = desc.partition("\n\n")
            txt.insert("end", "\n" + hint + "\n")
            if info:
                txt.insert("end", "\n" + info, "i")
        txt.config(state="disabled")
        self._fill_where(self.it_where, ITEM, iid)

    # --- miejsca występowania ----------------------------------------------------

    def _fill_where(self, tv, kind, iid, kinds=None):
        tv.delete(*tv.get_children())
        self._where_keys = getattr(self, "_where_keys", {})
        rows = []
        for k in (kinds or [kind]):
            rows += self.world.by_id.get((k, iid), [])
        for n, (p, e) in enumerate(rows):
            note = ROLA.get(e.role, e.role) + (f" ({e.note})" if e.note and e.kind != ITEM else "")
            if e.kind == ARMY and e.role != "bohater":
                note = "w armii wroga"
            key = f"w{n}"
            tv.insert("", "end", iid=key, values=(self.map_label(p), self.place_label(p), note, e.count,
                                                  self.map_state(p)))
            self._where_keys[(str(tv), key)] = p.key

    def _goto_where(self, tv):
        sel = tv.selection()
        if sel:
            key = self._where_keys.get((str(tv), sel[0]))
            if key:
                self.goto_place(key)

    def show_where(self, tree, tv, kind):
        sel = tree.selection()
        if not sel:
            tv.delete(*tv.get_children())
            return
        kinds = [UNIT, ARMY] if kind == UNIT else [kind]
        self._fill_where(tv, kind, sel[0], kinds)

    # --- zakładka 3 ------------------------------------------------------------

    def fill_spells(self):
        if not self.world:
            return
        t = self.sp_tree
        t.delete(*t.get_children())
        g, w = self.game, self.world
        q = self.sp_q.get().strip().casefold()
        school = self.sp_school.get()
        book, scrolls = w.hero["czary_w_ksiedze"], w.hero["zwoje"]
        ids = {i for i, s in g.spells.items()
               if s["school"] in SZKOLY and s["file"] in ("spells.txt", "spells_adventure.txt")
               and (s.get("image") or i in book or i in scrolls)}
        ids |= {i for (k, i) in w.by_id if k == SPELL}
        n = 0
        for sid in ids:
            sp = g.spells.get(sid) or {}
            sch = SZKOLY.get(sp.get("school"), "")
            if school not in ("", "wszystkie") and sch != school:
                continue
            name = g.spell_name(sid)
            if q and q not in name.casefold() and q not in sid.casefold():
                continue
            where = [x for x in w.by_id.get((SPELL, sid), []) if x[0].source != "bohater"]
            shops = sum(1 for p, e in where if e.role == "sklep")
            if sid in book:
                status, tag = f"Znasz (poziom {book[sid]})", "masz"
            elif where:
                status, tag = f"Do zdobycia: {len(where)} {miejsc(len(where))}" + \
                              (f", w tym {shops} na sprzedaż" if shops else ""), "jest"
            elif sid in scrolls:
                status, tag = "Masz tylko zwoje", "masz"
            else:
                status, tag = "Nie ma go w świecie gry (może dać zadanie lub rozmowa)", "brak"
            t.insert("", "end", iid=sid, tags=(tag,), **self.icon_kw(SPELL, sid), values=(
                name, sch, sp.get("level", ""), sp.get("price") or "", book.get(sid, ""),
                scrolls.get(sid, ""), len(where) or "", status))
            n += 1
        for i, k in enumerate(sorted(t.get_children(), key=lambda k: t.set(k, "name").casefold())):
            t.move(k, "", i)
        self.nb.tab(2, text=f"Czary ({n})")
        self.show_spell()

    # --- zakładka 4 ------------------------------------------------------------

    def fill_units(self):
        if not self.world:
            return
        t = self.un_tree
        t.delete(*t.get_children())
        w = self.world
        q = self.un_q.get().strip().casefold()
        mine = {}
        for a in w.hero["armia"]:
            mine[a["jednostka"]] = mine.get(a["jednostka"], 0) + a["ilosc"]
        ids = {i for (k, i) in w.by_id if k in (UNIT, ARMY)}
        n = 0
        for uid in ids:
            rec = [x for x in w.by_id.get((UNIT, uid), [])]
            enemy = [x for x in w.by_id.get((ARMY, uid), []) if x[0].source != "bohater"]
            shops = [x for x in rec if x[1].role in ("werbunek", "garnizon")]
            join = [x for x in rec if x[1].role == "dolaczaja"]
            if not shops and not join and not self.un_enemy.get():
                continue
            name = self.game.unit_name(uid)
            if q and q not in name.casefold() and q not in uid.casefold():
                continue
            total = sum(e.count for _, e in shops if isinstance(e.count, int) and e.role == "werbunek")
            t.insert("", "end", iid=uid, tags=("jest",) if shops else ("brak",), **self.icon_kw(UNIT, uid), values=(
                name, len(shops) or "", total or "", sum(e.count for _, e in join if isinstance(e.count, int)) or "",
                len(enemy) or "", mine.get(uid, "")))
            n += 1
        for i, k in enumerate(sorted(t.get_children(), key=lambda k: t.set(k, "name").casefold())):
            t.move(k, "", i)
        self.nb.tab(3, text=f"Jednostki ({n})")
        self.show_unit()

    # --- karty czaru i jednostki --------------------------------------------------

    def show_spell(self):
        sel = self.sp_tree.selection()
        txt = self.sp_desc
        self._text_begin(txt)
        if not sel:
            self.sp_where.delete(*self.sp_where.get_children())
            txt.insert("end", "Wybierz czar z listy.", "h")
            txt.config(state="disabled")
            return
        sid, g, w = sel[0], self.game, self.world
        sp = g.spells.get(sid) or {}
        self._text_image(txt, self.photo(SPELL, sid, (183, 104), big=True), self.photo(SPELL, sid, (90, 104)))
        txt.insert("end", g.spell_name(sid) + "\n", "h")
        facts = [SZKOLY.get(sp.get("school"), "")]
        if sp.get("level"):
            facts.append(f"poziom {sp['level']}")
        if sp.get("price"):
            facts.append(f"cena zwoju {sp['price']}")
        txt.insert("end", ", ".join(f for f in facts if f) + f"\nid: {sid}\n", "s")
        status = self.sp_tree.set(sid, "status")
        txt.insert("end", "\n" + status + "\n", self.sp_tree.item(sid, "tags")[0])
        d = g.spell_description(sid)
        if d["small"]:
            txt.insert("end", "\n" + d["small"] + "\n", "b")
        if d["desc"]:
            txt.insert("end", "\n" + d["desc"] + "\n")
        if d["levels"]:
            txt.insert("end", "Poziomy czaru\n", "h2")
            for lvl, mana, crystals, effect in d["levels"]:
                cost = ", ".join(x for x in (f"{mana} many" if mana else "",
                                             f"{crystals} kryształów" if crystals else "") if x)
                txt.insert("end", f"Poziom {lvl}", "b")
                txt.insert("end", f" ({cost})" if cost else "")
                txt.insert("end", (f": {effect}" if effect else "") + "\n", "li")
        params = {k: v for k, v in (sp.get("params") or {}).items() if v}
        if params:
            txt.insert("end", "Parametry z pliku gry\n", "h2")
            txt.insert("end", ", ".join(f"{k} = {v}" for k, v in params.items()) + "\n", "s")
            txt.insert("end", "Liczby w opisie („…”) gra wylicza w trakcie gry, m.in. z intelektu bohatera.\n", "s")
        txt.config(state="disabled")
        self._fill_where(self.sp_where, SPELL, sid)

    def show_unit(self):
        sel = self.un_tree.selection()
        txt = self.un_desc
        self._text_begin(txt)
        if not sel:
            self.un_where.delete(*self.un_where.get_children())
            txt.insert("end", "Wybierz jednostkę z listy.", "h")
            txt.config(state="disabled")
            return
        uid, g = sel[0], self.game
        self._text_image(txt, self.photo(UNIT, uid, (120, 153), big=True))
        txt.insert("end", g.unit_name(uid) + "\n", "h")
        u = g.unit_info(uid)
        if not u:
            txt.insert("end", f"id: {uid}\nBrak danych jednostki w plikach gry.", "s")
        else:
            facts = [RASY.get(u["race"], u["race"]), f"poziom {u['level']}" if u["level"] else "",
                     f"przywództwo {u['leadership']}" if u["leadership"] else "",
                     f"koszt {u['cost']} zł" if u["cost"] else ""]
            txt.insert("end", ", ".join(f for f in facts if f) + f"\nid: {uid}\n", "s")
            stats = [("Atak", u["attack"]), ("Obrona", u["defense"]), ("Zdrowie", u["hitpoint"]),
                     ("Szybkość", u["speed"]), ("Inicjatywa", u["initiative"]),
                     ("Trafienie krytyczne", f"{u['krit']}%" if u["krit"] else "")]
            txt.insert("end", "\n")
            for i, (label, val) in enumerate(x for x in stats if x[1]):
                txt.insert("end", ("  •  " if i else "") + f"{label} ")
                txt.insert("end", str(val), "b")
            txt.insert("end", "\n")
            if u["resistances"]:
                txt.insert("end", "Odporności: " + ", ".join(f"{k} {v}%" for k, v in u["resistances"].items())
                           + "\n", "s")
            if u["features"]:
                txt.insert("end", "Cechy\n", "h2")
                for title, hint in u["features"]:
                    txt.insert("end", "• ", "li")
                    txt.insert("end", title, ("b", "li"))
                    txt.insert("end", (f" – {hint}" if hint else "") + "\n", "li")
            if u["attacks"]:
                txt.insert("end", "Ataki i zdolności\n", "h2")
                for a in u["attacks"]:
                    txt.insert("end", "• ", "li")
                    txt.insert("end", a["name"], ("b", "li"))
                    extra = " – ".join(x for x in (f"obrażenia {a['damage']}" if a["damage"] else "", a["hint"]) if x)
                    txt.insert("end", (f": {extra}" if extra else "") + "\n", "li")
        mine = self.un_tree.set(uid, "mine")
        if mine:
            txt.insert("end", f"\nW twojej armii: {mine}\n", "masz")
        txt.config(state="disabled")
        self._fill_where(self.un_where, UNIT, uid, [UNIT, ARMY])

    # --- zakładka 5 ------------------------------------------------------------

    def fill_hero(self):
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            print_report(self.world.hero)
        text = buf.getvalue().strip("\n")
        g = self.game
        # identyfikatory przedmiotów, czarów i jednostek -> nazwy z gry
        lines = []
        for line in text.splitlines():
            names = []
            for token in line.replace(",", " ").split():
                name = None
                if token in g.items:
                    name = g.item_name(token)
                elif token in g.spells:
                    name = g.spell_name(token)
                elif g.text(f"cpn_{token}"):
                    name = g.unit_name(token)
                if name and name != token:
                    names.append(name)
            lines.append(f"{line:<48}  {', '.join(names)}" if names else line)
        self.hero_txt.config(state="normal")
        self.hero_txt.delete("1.0", "end")
        self.hero_txt.insert("end", "\n".join(lines))
        self.hero_txt.config(state="disabled")

    # --- eksport ---------------------------------------------------------------

    def export(self):
        if not self.world:
            return
        path = filedialog.asksaveasfilename(
            title="Zapisz listę", defaultextension=".xlsx",
            initialfile=Path(self.save_path).stem + "_swiat.xlsx",
            filetypes=[("Skoroszyt Excela", "*.xlsx"), ("CSV (średniki)", "*.csv")])
        if not path:
            return
        try:
            export_world(self, path)
        except ImportError:
            messagebox.showerror(APP_TITLE, "Do zapisu .xlsx potrzebny jest openpyxl (pip install openpyxl).\n"
                                            "Możesz wybrać format .csv.")
            return
        except OSError as e:
            messagebox.showerror(APP_TITLE, f"Nie udało się zapisać pliku:\n{e}")
            return
        self.status.config(text=f"Zapisano: {path}")


SHORT = {ITEM: "przedm.", SPELL: "czary", UNIT: "jedn.", ARMY: "w armii", RES: "zasoby"}
RES_ICONS = {"crystals": "Message_icon1_crystals.png", "rune_might": "Message_icon1_rune_might.png",
             "rune_mind": "Message_icon1_rune_mind.png", "rune_magic": "Message_icon1_rune_magic.png",
             "money": "message_icon1_money.png", "leadership": "message_icon1_leadership.png",
             "mana": "message_icon1_mana.png", "rage": "message_icon1_rage.png"}


def export_world(app, path):
    w, g = app.world, app.game
    places = [["Mapa", "Stan mapy", "Miejsce", "Rodzaj miejsca", "Kategoria", "Nazwa", "ID", "Ilość",
               "Typ / szkoła", "Poziom", "Jak zdobyć", "Uwagi"]]
    for p in w.places:
        for e in p.entries:
            typ, lvl = app.entry_type(e)
            places.append([app.map_label(p), app.map_state(p), app.place_label(p), p.type, e.kind,
                           app.ename(e.kind, e.id), e.id, e.count, typ, lvl, ROLA.get(e.role, e.role), e.note])
    items = [["Nazwa", "ID", "Typ", "Premie", "Poziom", "Rasa", "Cena", "Limit w grze", "Wylosowano", "Miejsc",
              "Status"]]
    for iid in sorted(app._item_ids(), key=lambda i: g.item_name(i).casefold()):
        it = g.items.get(iid) or {}
        if it.get("slot") in HIDDEN_SLOTS:
            continue
        where = len([x for x in w.by_id.get((ITEM, iid), []) if x[0].source != "bohater"])
        items.append([g.item_name(iid), iid, app.item_type(iid),
                      ", ".join(bonus_text(bn) for bn in it.get("bonuses", [])), it.get("level", ""),
                      RASY.get(it.get("race"), it.get("race", "")), it.get("price"), it.get("maxcount"),
                      w.box.get(iid, 0), where, w.item_status(iid)[1]])
    path = Path(path)
    if path.suffix.lower() == ".csv":
        import csv
        for p, rows in ((path, places), (path.with_name(path.stem + "_przedmioty.csv"), items)):
            with p.open("w", newline="", encoding="utf-8-sig") as f:
                csv.writer(f, delimiter=";").writerows(rows)
        return
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    for ws, rows, widths in ((wb.active, places, [26, 30, 40, 18, 16, 30, 24, 8, 16, 8, 18, 28]),
                             (wb.create_sheet(), items, [30, 24, 16, 50, 8, 12, 9, 12, 12, 9, 60])):
        ws.title = "Miejsca" if rows is places else "Przedmioty"
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", start_color="305496")
        for i, wdt in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = wdt
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    try:
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except (AttributeError, OSError, ImportError):
        pass
    App(sys.argv[1] if len(sys.argv) > 1 else None).mainloop()


if __name__ == "__main__":
    main()
