"""Odczyt zapisów gry King's Bounty: Armored Princess (.sav).

Plik .sav to archiwum ZIP z trzema częściami:
    saveinfo  - podsumowanie slotu (poziom, armia, zrzut ekranu, miniatura)
    savedata  - pełny stan gry, nagłówek 'slcb' + strumień zlib
    savecrc   - suma kontrolna (algorytm nieznany, skrypt tylko czyta)

Obie części danych używają tego samego formatu drzewa:
    węzeł  := u32 rodzaj, [klucz], ciało
    ciało  := u32 n, n * u32 typ, n * pole
    pole   := klucz (u32 dł. + bajty) + wartość zależna od typu
Typy: 0 blob, 1 węzeł, 2 tekst UTF-16, 3 tekst, 4 int32, 5 uint32,
      6 float32, 7 float64. Dzieci węzłów rodzaju 6 i 7 nie mają kluczy.

Format ustalony na podstawie analizy zapisów, a nie dokumentacji,
więc znaczenie części pól jest domysłem (zaznaczonym w raporcie).

Użycie:
    python kb_save.py Montero.sav
    python kb_save.py Montero.sav --json raport.json --obrazy obrazy
    python kb_save.py Montero.sav --zrzut drzewo.json
    python kb_save.py Montero.sav --lista Montero_lista.xlsx   (lub .csv)
"""

import argparse
import json
import struct
import sys
import zipfile
import zlib
from pathlib import Path

KEYLESS = {6, 7}                                   # rodzaje węzłów-tablic
KLASY = {0: "Wojownik", 1: "Paladyn", 2: "Mag"}    # założenie na podstawie statystyk


# --- Parser drzewa -----------------------------------------------------------

class Node(list):
    """Węzeł drzewa: lista par (klucz, wartość) + rodzaj węzła."""

    def __init__(self, kind, items):
        super().__init__(items)
        self.kind = kind

    def get(self, key, default=None):
        for k, v in self:
            if k == key:
                return v
        return default

    def path(self, *keys):
        node = self
        for k in keys:
            node = node.get(k) if isinstance(node, Node) else None
            if node is None:
                return None
        return node


class Reader:
    def __init__(self, data: bytes):
        self.d, self.p = data, 0

    def u32(self):
        v = struct.unpack_from("<I", self.d, self.p)[0]
        self.p += 4
        return v

    def raw(self):
        n = self.u32()
        v = self.d[self.p:self.p + n]
        self.p += n
        return v

    def body(self, kind):
        n = self.u32()
        types = [self.u32() for _ in range(n)]
        items = []
        for i, t in enumerate(types):
            if t == 1:
                sub_kind = self.u32()
                key = str(i) if kind in KEYLESS else self.raw().decode("latin1")
                items.append((key, self.body(sub_kind)))
                continue
            key = self.raw().decode("latin1")
            if t == 0:
                v = self.raw()
            elif t == 2:
                n2 = self.u32()
                v = self.d[self.p:self.p + 2 * n2].decode("utf-16le")
                self.p += 2 * n2
            elif t == 3:
                v = self.raw().decode("utf-8", "replace")
            elif t == 4:
                v = struct.unpack_from("<i", self.d, self.p)[0]
                self.p += 4
            elif t == 5:
                v = self.u32()
            elif t == 6:
                v = struct.unpack_from("<f", self.d, self.p)[0]
                self.p += 4
            elif t == 7:
                v = struct.unpack_from("<d", self.d, self.p)[0]
                self.p += 8
            else:
                raise ValueError(f"Nieznany typ pola {t} na pozycji {self.p}")
            items.append((key, v))
        return Node(kind, items)

    def root(self):
        kind = self.u32()
        self.raw()                       # nazwa korzenia
        node = self.body(kind)
        if self.p != len(self.d):
            raise ValueError(f"Nie odczytano całości: {self.p} z {len(self.d)} bajtów")
        return node


def load_save(path):
    with zipfile.ZipFile(path) as z:
        info_raw = z.read("saveinfo")
        data_raw = z.read("savedata")
    if data_raw[:4] != b"slcb":
        raise ValueError("Nieoczekiwany nagłówek savedata (brak 'slcb')")
    size, csize = struct.unpack_from("<II", data_raw, 4)
    data = zlib.decompress(data_raw[12:12 + csize])
    if len(data) != size:
        raise ValueError("Rozmiar po dekompresji nie zgadza się z nagłówkiem")
    return Reader(info_raw).root(), Reader(data).root()


# --- Wyciąganie danych bohatera ---------------------------------------------

def find_player_hero(data):
    """Bohater gracza: węzeł hero@N w lubodies z ekwipunkiem i armią."""
    bodies = data.path("session", "lus", "lubodies")
    for key, node in bodies or []:
        if key.startswith("hero@") and isinstance(node, Node):
            vars_ = node.get("vars")
            if vars_ is not None and vars_.get(".items") is not None:
                return key, vars_
    raise ValueError("Nie znaleziono bohatera gracza w zapisie")


def strg(vars_, key):
    node = vars_.get(key)
    return node.get("strg") if isinstance(node, Node) else None


def parse_lvars(text):
    parts = text.split("/") if text else []
    return {parts[i]: parts[i + 1] for i in range(0, len(parts) - 1, 2)}


def hero_report(info, data):
    hero_id, v = find_player_hero(data)
    summary = data.get("hero") or Node(2, [])

    stats, equipped, backpack, medals, other = {}, [], [], [], []
    for name, item in strg(v, ".items") or []:
        if not isinstance(item, Node):
            continue
        if item.get("count") is not None:             # zasób / statystyka
            stats[name] = {
                "bazowo": item.get("count"), "bonus": item.get("%count"),
                "limit": item.get("limit"), "bonus_limitu": item.get("%limit"),
            }
            continue
        lv = parse_lvars(item.get("lvars", ""))
        lv.pop("rndid", None)
        entry = {"przedmiot": name, **({"parametry": lv} if lv else {})}
        if item.get("slruck"):
            pos, _, cnt = item.get("slruck").partition(",")
            backpack.append({**entry, "pozycja": pos, "ilosc": int(cnt or 1)})
        elif str(item.get("slbody", "")).startswith("@"):
            (medals if name.startswith("medal_") else other).append(
                name if name.startswith("medal_") else {**entry, "slot": "ukryty"})
        elif item.get("slbody") == "wife":
            other.append({**entry, "slot": "towarzysz"})
        else:
            equipped.append({**entry, "slot": item.get("slbody")})

    def total(k):
        s = stats.get(k, {})
        return (s.get("bazowo") or 0) + (s.get("bonus") or 0)

    def maximum(k):
        s = stats.get(k, {})
        return (s.get("limit") or 0) + (s.get("bonus_limitu") or 0)

    stack = (strg(v, ".army") or Node(2, [])).get("stack", "")
    parts = stack.split("|") if stack else []
    army = [{"jednostka": parts[i], "ilosc": int(parts[i + 1])}
            for i in range(0, len(parts) - 1, 2)]

    spells = strg(v, ".spells") or Node(2, [])
    pet = (strg(v, ".spirits") or Node(2, [])).get("pet") or Node(2, [])
    pet_upgrades = {}
    for u in (pet.get("upgs") or "").split("/"):
        name, _, lvl = u.partition(".")
        if name:
            pet_upgrades.setdefault(name, []).append(lvl)

    visited = sorted(m for m, node in (data.path("session", "sesd", "maps") or [])
                     if isinstance(node, Node) and node.path("lts", hero_id) is not None)

    klasa = summary.get("c")
    return {
        "zapis": {
            "dzien_gry": info.get("game_day"),
            "godzina_gry": round(info.get("game_time") or 0, 2),
            "lokacja": info.get("location"),
            "poziom_trudnosci": summary.get("d"),
        },
        "bohater": {
            "imie": summary.get("n") or strg(v, "name"),
            "klasa": f"{klasa} ({KLASY.get(klasa, '?')})",
            "poziom": strg(v, "level"),
            "doswiadczenie": f"{stats.get('experience', {}).get('bazowo')} / "
                             f"{stats.get('experience', {}).get('limit')}",
            "atak": total("attack"), "obrona": total("defense"),
            "intelekt": total("intellect"), "przywodztwo": total("leadership"),
            "mana": f"{total('mana')} / {maximum('mana')}",
            "furia": f"{total('rage')} / {maximum('rage')}",
            "zloto": total("money"), "krysztaly": total("crystals"),
            "runy": {"sily": total("rune_might"), "umyslu": total("rune_mind"),
                     "magii": total("rune_magic")},
            "pojemnosc_ksiegi": total("booksize"),
        },
        "armia": army,
        "umiejetnosci": dict(strg(v, ".skills") or []),
        "czary_w_ksiedze": dict(spells.get("m") or []),
        "zwoje": dict(spells.get("s") or []),
        "zwierzak": {"imie": summary.get("pn"), "poziom": pet.get("lev"),
                     "doswiadczenie": f"{pet.get('spxp')} / {pet.get('nxp')}",
                     "ulepszenia": pet_upgrades},
        "zalozone": sorted(equipped, key=lambda e: int(e["slot"]) if str(e["slot"]).isdigit() else 99),
        "plecak": backpack,
        "towarzysz_i_ukryte": other,
        "medale": sorted(medals),
        "premie": {k: s["bonus"] for k, s in stats.items() if k.startswith("sp_")},
        "mapy_z_wizyta_bohatera": visited,
    }


# --- Wyjście ------------------------------------------------------------------

def print_report(r):
    b, z = r["bohater"], r["zapis"]
    print(f"\n=== {b['imie']} | poziom {b['poziom']} | {b['klasa']} ===")
    print(f"Dzień {z['dzien_gry']}, godz. {z['godzina_gry']}, lokacja: {z['lokacja']}, "
          f"trudność: {z['poziom_trudnosci']}")
    print(f"Doświadczenie: {b['doswiadczenie']}")
    print(f"Atak {b['atak']} | Obrona {b['obrona']} | Intelekt {b['intelekt']} | "
          f"Przywództwo {b['przywodztwo']}")
    print(f"Mana {b['mana']} | Furia {b['furia']} | Złoto {b['zloto']} | "
          f"Kryształy {b['krysztaly']}")
    runy = b["runy"]
    print(f"Runy: siły {runy['sily']}, umysłu {runy['umyslu']}, magii {runy['magii']}")

    print("\nArmia (pozycje powyżej 5 to prawdopodobnie rezerwa):")
    for i, a in enumerate(r["armia"], 1):
        print(f"  {i}. {a['jednostka']:<14} {a['ilosc']:>5}")

    def section(title, d):
        print(f"\n{title}:")
        for k, val in sorted(d.items(), key=lambda kv: (-kv[1] if isinstance(kv[1], int) else 0, kv[0])):
            print(f"  {k:<28} {val}")

    section("Umiejętności (-1 = prawdopodobnie niewykupiona)", r["umiejetnosci"])
    section("Czary w księdze (poziom)", r["czary_w_ksiedze"])
    section("Zwoje (ilość)", r["zwoje"])

    p = r["zwierzak"]
    print(f"\nZwierzak: {p['imie']}, poziom {p['poziom']}, doświadczenie {p['doswiadczenie']}")
    for name, lv in p["ulepszenia"].items():
        print(f"  {name:<14} {', '.join(lv)}")

    print("\nZałożone przedmioty:")
    for e in r["zalozone"]:
        print(f"  slot {str(e['slot']):>3}: {e['przedmiot']}")
    print("\nPlecak:")
    for e in r["plecak"]:
        print(f"  {e['przedmiot']:<22} x{e['ilosc']}")
    for t in r["towarzysz_i_ukryte"]:
        print(f"\n{t['slot'].capitalize()}: {t['przedmiot']} {t.get('parametry', '')}")
    print(f"\nMedale: {', '.join(r['medale'])}")
    print(f"Mapy z wizytą bohatera: {', '.join(r['mapy_z_wizyta_bohatera'])}")


# --- Lista świata: obiekty na mapach, nagrody za zadania, statystyki ---------

TYPY_OBIEKTOW = {
    "building_trader": "Budynek / sklep", "building_castle": "Zamek",
    "template_item_altar": "Ołtarz / skrytka", "template_item_mb": "Skrzynia / znalezisko",
    "template_item_ft": "Bonus na mapie", "template_mana_fountain": "Źródło many",
    "army_a_template": "Armia", "army_n_template": "Armia", "NPC_template": "NPC",
    "map": "Mapa wyspy", "hero": "Bohater", "template_rage_fountain": "Źródło furii",
    "template_with_altar": "Ołtarz",
}
ALIGN = {"enemy": "wroga", "neutral": "neutralna"}
ZASOBY = {"money", "crystals", "leadership", "experience", "rune_might",
          "rune_mind", "rune_magic", "mana", "rage", "attack", "defense", "intellect"}
NAGRODY = {"bp": "Zasób", "bi": "Przedmiot", "bs": "Zwój / czar", "bu": "Jednostki"}


def _groups(text, sep, step=2):
    parts = (text or "").split(sep) if isinstance(text, str) else []
    for i in range(0, len(parts) - step + 1, step):
        if parts[i]:
            yield parts[i:i + step]


def world_rows(data, hero_id):
    """Zawartość obiektów na mapach: sklepy, zamki, skrzynie, armie."""
    rows = []
    for key, body in data.path("session", "lus", "lubodies") or []:
        if not isinstance(body, Node) or not body.get("lt") or key == hero_id:
            continue
        v = body.get("vars") or Node(2, [])
        typ = key.split("@")[0]
        nazwa_typu = TYPY_OBIEKTOW.get(typ, typ)
        if typ.startswith("army_"):
            align = strg(v, "align")
            nazwa_typu = f"Armia ({ALIGN.get(align, align)})" if align else "Armia"
        etykieta = strg(v, "hint") or ""
        if isinstance(etykieta, str):
            etykieta = etykieta.replace("itext_", "").replace("_hint", "")
        base = {"mapa": body.get("lt"), "obiekt": nazwa_typu, "id": key, "etykieta": etykieta}

        def add(kat, nazwa, ilosc, szczegoly=""):
            rows.append({**base, "kategoria": kat, "nazwa": nazwa,
                         "ilosc": ilosc, "szczegoly": szczegoly})

        for name, item in strg(v, ".items") or []:
            if not isinstance(item, Node):
                continue
            if item.get("count") is not None:
                add("Zasób" if name in ZASOBY else "Przedmiot", name, item.get("count"))
            else:
                _, _, cnt = (item.get("slruck") or "").partition(",")
                lv = parse_lvars(item.get("lvars", ""))
                lv.pop("rndid", None)
                add("Przedmiot", name, int(cnt or 1),
                    ", ".join(f"{a}={b}" for a, b in lv.items()))
        for unit, cnt in _groups(strg(v, ".shopunits"), "/"):
            add("Jednostki do werbunku", unit, int(cnt))
        for unit, a, b in _groups(strg(v, ".castleunits"), "/", 3):
            add("Jednostka zamku", unit, int(a), f"druga wartość: {b}")
        spells = strg(v, ".spells")
        if isinstance(spells, Node):
            for kind, label in (("s", "Zwój na sprzedaż"), ("m", "Czar")):
                for name, cnt in spells.get(kind) or []:
                    add(label, name, cnt)
        army = strg(v, ".army")
        stack = army.get("stack") if isinstance(army, Node) else ""
        hero_face = strg(v, ".face")
        for unit, cnt in _groups(stack, "|"):
            if cnt:
                add("Armia (obrońcy / wrogowie)", unit, int(cnt),
                    f"z bohaterem: {hero_face}" if hero_face else "")
    rows.sort(key=lambda r: (r["mapa"], r["obiekt"], r["id"], r["kategoria"], r["nazwa"]))
    return rows


def quest_reward_rows(data):
    """Nagrody za zadania wylosowane z góry (generatory powiązane z zadaniami)."""
    status = {str(q.get("uid")): q.get("sta") for _, q in data.path("server", "quest") or []
              if isinstance(q, Node)}
    quests = {k.split("_", 1)[1] for k, _ in data.path("session", "lus", "lubodies") or []
              if k.startswith("quest_")}
    rows = []

    def walk(node, qid):
        if not isinstance(node, Node):
            return
        sn = node.get("sn")
        if sn in NAGRODY:
            nazwa = node.get("ib") or node.get("uni") or ""
            ilosc = node.get("va") if node.get("va") is not None else node.get("ko")
            if sn == "bu" and "/" in nazwa:
                nazwa, _, ilosc = nazwa.partition("/")
            rows.append({"zadanie": qid, "status_zadania": status.get(qid, ""),
                         "typ": NAGRODY[sn], "nazwa": nazwa, "ilosc": ilosc})
        for _, sub in node:
            walk(sub, qid)

    for _, el in data.path("server", "embs", "els") or []:
        qid = (el.get("n") or "").strip("_").split("_")[0]
        if qid in quests:
            walk(el.get("pregen"), qid)
    rows.sort(key=lambda r: (r["zadanie"], r["typ"], r["nazwa"]))
    return rows


def game_stats(data):
    gv = data.get("gv") or Node(2, [])
    return [[k, v] for k, v in gv if isinstance(v, (int, float))]


def export_list(path, rows, quests, stats):
    path = Path(path)
    obj_cols = [("mapa", "Mapa"), ("obiekt", "Obiekt"), ("id", "ID obiektu"),
                ("etykieta", "Etykieta"), ("kategoria", "Kategoria"), ("nazwa", "Nazwa"),
                ("ilosc", "Ilość"), ("szczegoly", "Szczegóły")]
    q_cols = [("zadanie", "ID zadania"), ("status_zadania", "Status (surowy)"),
              ("typ", "Typ nagrody"), ("nazwa", "Nazwa"), ("ilosc", "Ilość / wartość")]

    if path.suffix.lower() == ".csv":
        import csv
        for p, cols, data_rows in ((path, obj_cols, rows),
                                   (path.with_name(path.stem + "_zadania.csv"), q_cols, quests)):
            with p.open("w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f, delimiter=";")
                w.writerow([h for _, h in cols])
                for r in data_rows:
                    w.writerow([r[k] for k, _ in cols])
            print(f"Zapisano: {p}")
        return

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("Do eksportu .xlsx potrzebny jest openpyxl: pip install openpyxl "
                 "(albo podaj plik .csv)")

    font = Font(name="Arial", size=10)
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", start_color="305496")

    def style_header(ws):
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    def sheet(ws, headers, data_rows, widths):
        ws.append(headers)
        for row in data_rows:
            ws.append(row)
        style_header(ws)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = font
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    wb = Workbook()
    ws = wb.active
    ws.title = "Obiekty"
    sheet(ws, [h for _, h in obj_cols], [[r[k] for k, _ in obj_cols] for r in rows],
          [22, 24, 28, 26, 26, 28, 9, 44])
    n = len(rows) + 1

    wq = wb.create_sheet("Nagrody za zadania")
    sheet(wq, [h for _, h in q_cols], [[r[k] for k, _ in q_cols] for r in quests],
          [14, 16, 16, 30, 16])

    wst = wb.create_sheet("Statystyki")
    sheet(wst, ["Statystyka", "Wartość"], stats, [32, 14])

    wsum = wb.create_sheet("Podsumowanie", 0)
    cats = ["Przedmiot", "Zasób", "Jednostki do werbunku", "Zwój na sprzedaż",
            "Armia (obrońcy / wrogowie)", "Jednostka zamku"]
    maps = sorted({r["mapa"] for r in rows})
    wsum.append(["Mapa", "Wierszy razem"] + cats)
    for i, m in enumerate(maps, 2):
        wsum.append([m, f"=COUNTIF(Obiekty!$A$2:$A${n},$A{i})"] +
                    [f'=COUNTIFS(Obiekty!$A$2:$A${n},$A{i},Obiekty!$E$2:$E${n},"{c}")'
                     for c in cats])
    last = len(maps) + 1
    wsum.append(["RAZEM"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{last})"
                             for c in range(2, len(cats) + 3)])
    style_header(wsum)
    for row in wsum.iter_rows(min_row=2):
        for cell in row:
            cell.font = font
    for cell in wsum[last + 1]:
        cell.font = Font(name="Arial", size=10, bold=True)
    wsum.column_dimensions["A"].width = 26
    for c in range(2, len(cats) + 3):
        wsum.column_dimensions[get_column_letter(c)].width = 15
    wsum.row_dimensions[1].height = 30
    wsum.freeze_panes = "B2"
    note = last + 3
    wsum.cell(note, 1, "Uwagi").font = Font(name="Arial", size=10, bold=True)
    for j, t in enumerate([
        "Lista obejmuje tylko odwiedzone mapy: gra losuje zawartość wyspy przy pierwszej wizycie.",
        "Nazwy to identyfikatory wewnętrzne gry. Ekwipunek bohatera gracza jest pominięty.",
        "Arkusz 'Nagrody za zadania': nagrody wylosowane z góry; część może dotyczyć zadań już ukończonych.",
        "Znaczenie pola 'Status (surowy)' i drugiej wartości przy jednostkach zamku nie jest potwierdzone.",
    ], 1):
        wsum.cell(note + j, 1, t).font = font

    wb.save(path)
    print(f"Lista zapisana: {path}")


def save_images(info, out_dir):
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    for key in ("screenshot", "slotpicture"):
        blob = info.get(key)
        if isinstance(blob, bytes) and blob.startswith(b"\x89PNG"):
            (out / f"{key}.png").write_bytes(blob)
            print(f"Zapisano {out / (key + '.png')}")


def to_jsonable(node):
    if isinstance(node, Node):
        return {"_rodzaj": node.kind, "_pola": [[k, to_jsonable(v)] for k, v in node]}
    if isinstance(node, bytes):
        return {"_bajty": len(node), "_hex": node[:64].hex()}
    return node


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="Odczyt zapisu King's Bounty: Armored Princess")
    ap.add_argument("plik", help="ścieżka do pliku .sav")
    ap.add_argument("--json", help="zapisz raport do pliku JSON")
    ap.add_argument("--obrazy", help="katalog na zrzut ekranu i miniaturę z zapisu")
    ap.add_argument("--zrzut", help="zapisz całe drzewo savedata do JSON (do eksploracji)")
    ap.add_argument("--lista", help="lista sklepów, zamków, skrzyń, armii i nagród "
                                    "za zadania: plik .xlsx (wymaga openpyxl) lub .csv")
    args = ap.parse_args()

    info, data = load_save(args.plik)
    report = hero_report(info, data)
    print_report(report)

    if args.json:
        Path(args.json).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"\nRaport JSON: {args.json}")
    if args.obrazy:
        save_images(info, args.obrazy)
    if args.lista:
        hero_id, _ = find_player_hero(data)
        export_list(args.lista, world_rows(data, hero_id), quest_reward_rows(data),
                    game_stats(data))
    if args.zrzut:
        Path(args.zrzut).write_text(json.dumps(to_jsonable(data), ensure_ascii=False), encoding="utf-8")
        print(f"Drzewo zapisu: {args.zrzut}")


if __name__ == "__main__":
    main()
